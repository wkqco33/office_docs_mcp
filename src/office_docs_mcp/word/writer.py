from __future__ import annotations

from docx import Document

from office_docs_mcp.common.file_utils import (
    create_backup,
    ensure_parent_dir,
    validate_file_path,
)

VALID_WORD_EXTS = (".docx",)


def word_create_document(file_path: str, title: str | None = None) -> str:
    """Create a new empty Word document with an optional title.

    Args:
        file_path: Path where the document will be saved (.docx).
        title: Optional document title to add as the first heading.

    Returns:
        Status confirmation message.
    """
    path = ensure_parent_dir(file_path)
    if path.suffix.lower() not in VALID_WORD_EXTS:
        path = path.with_suffix(".docx")

    doc = Document()
    if title:
        doc.add_heading(title, level=0)

    doc.save(str(path))
    return f"Created Word document at {path}"


def word_append_paragraph(
    file_path: str,
    text: str,
    style: str | None = None,
    backup: bool = False,
) -> str:
    """Append a paragraph or heading to the end of a Word document.

    Args:
        file_path: Path to the Word document (.docx).
        text: Paragraph text content.
        style: Optional style name (e.g., 'Heading 1', 'Heading 2', 'List Bullet').
        backup: If True, create a backup before modifying.

    Returns:
        Status confirmation message.
    """
    path = validate_file_path(file_path, expected_extensions=VALID_WORD_EXTS)
    bak_msg = ""
    if backup:
        bak = create_backup(path)
        if bak:
            bak_msg = f" (backup created: {bak.name})"

    doc = Document(str(path))

    try:
        if style and style.lower().startswith("heading"):
            # Extract heading level if possible
            try:
                level = int(style.split()[-1])
                doc.add_heading(text, level=level)
            except (ValueError, IndexError):
                doc.add_paragraph(text, style=style)
        else:
            doc.add_paragraph(text, style=style)
    except KeyError as err:
        raise ValueError(f"Style '{style}' not found in document styles.") from err

    doc.save(str(path))
    return f"Appended paragraph to {path}{bak_msg}"


def word_append_table_row(
    file_path: str,
    table_idx: int,
    row_data: list[str],
    backup: bool = False,
) -> str:
    """Append a row of cells to an existing table in a Word document.

    Args:
        file_path: Path to the Word document (.docx).
        table_idx: 0-based index of the table.
        row_data: List of strings for each column cell in the new row.
        backup: If True, create a backup before modifying.

    Returns:
        Status confirmation message.
    """
    path = validate_file_path(file_path, expected_extensions=VALID_WORD_EXTS)
    bak_msg = ""
    if backup:
        bak = create_backup(path)
        if bak:
            bak_msg = f" (backup created: {bak.name})"

    doc = Document(str(path))

    if table_idx < 0 or table_idx >= len(doc.tables):
        raise IndexError(
            f"Table index {table_idx} out of range. Document has {len(doc.tables)} tables."
        )

    table = doc.tables[table_idx]
    row = table.add_row()
    for c_idx, val in enumerate(row_data):
        if c_idx < len(row.cells):
            row.cells[c_idx].text = str(val)

    doc.save(str(path))
    return f"Appended row to table {table_idx} in {path}{bak_msg}"


def word_write_table_cell(
    file_path: str,
    table_idx: int,
    row_idx: int,
    col_idx: int,
    text: str,
    backup: bool = False,
) -> str:
    """Update text in a specific cell of a table in a Word document.

    Args:
        file_path: Path to the Word document (.docx).
        table_idx: 0-based index of the table.
        row_idx: 0-based row index.
        col_idx: 0-based column index.
        text: New text content for the cell.
        backup: If True, create a backup before modifying.

    Returns:
        Status confirmation message.
    """
    path = validate_file_path(file_path, expected_extensions=VALID_WORD_EXTS)
    bak_msg = ""
    if backup:
        bak = create_backup(path)
        if bak:
            bak_msg = f" (backup created: {bak.name})"

    doc = Document(str(path))

    if table_idx < 0 or table_idx >= len(doc.tables):
        raise IndexError(
            f"Table index {table_idx} out of range. Document has {len(doc.tables)} tables."
        )

    table = doc.tables[table_idx]
    if row_idx < 0 or row_idx >= len(table.rows):
        raise IndexError(f"Row index {row_idx} out of range in table {table_idx}.")

    row = table.rows[row_idx]
    if col_idx < 0 or col_idx >= len(row.cells):
        raise IndexError(f"Column index {col_idx} out of range in table {table_idx}.")

    row.cells[col_idx].text = text
    doc.save(str(path))
    return f"Updated cell [{row_idx}, {col_idx}] of table {table_idx} in {path}{bak_msg}"


def word_replace_text(
    file_path: str,
    find_text: str,
    replace_text: str,
    count: int = -1,
    backup: bool = False,
) -> str:
    """Find and replace text across paragraphs and tables in a Word document.

    Args:
        file_path: Path to the Word document (.docx).
        find_text: Text string to search for.
        replace_text: Text to replace with.
        count: Max replacements to perform (-1 for unlimited).
        backup: If True, create a backup before modifying.

    Returns:
        Status confirmation message with replacement count.
    """
    if not find_text:
        raise ValueError("find_text cannot be empty.")

    path = validate_file_path(file_path, expected_extensions=VALID_WORD_EXTS)
    bak_msg = ""
    if backup:
        bak = create_backup(path)
        if bak:
            bak_msg = f" (backup created: {bak.name})"

    doc = Document(str(path))
    replacements = 0

    # 1. Replace in paragraphs
    for p in doc.paragraphs:
        if find_text in p.text:
            occurrences = p.text.count(find_text)
            if count != -1 and (replacements + occurrences) > count:
                max_in_p = max(0, count - replacements)
                p.text = p.text.replace(find_text, replace_text, max_in_p)
                replacements += max_in_p
                break
            p.text = p.text.replace(find_text, replace_text)
            replacements += occurrences
            if count != -1 and replacements >= count:
                break

    # 2. Replace in tables if count not reached
    if count == -1 or replacements < count:
        for table in doc.tables:
            for row in table.rows:
                for cell in row.cells:
                    if find_text in cell.text:
                        occurrences = cell.text.count(find_text)
                        if count != -1 and (replacements + occurrences) > count:
                            max_in_c = max(0, count - replacements)
                            cell.text = cell.text.replace(find_text, replace_text, max_in_c)
                            replacements += max_in_c
                            break
                        cell.text = cell.text.replace(find_text, replace_text)
                        replacements += occurrences
                        if count != -1 and replacements >= count:
                            break

    doc.save(str(path))
    return f"Replaced {replacements} occurrence(s) of '{find_text}' with '{replace_text}' in {path}{bak_msg}"


def word_delete_paragraph(
    file_path: str,
    paragraph_idx: int,
    backup: bool = False,
) -> str:
    """Delete a specific paragraph from a Word document by 0-based index.

    Args:
        file_path: Path to the Word document (.docx).
        paragraph_idx: 0-based index of the paragraph to delete.
        backup: If True, create a backup before modifying.

    Returns:
        Status confirmation message.
    """
    path = validate_file_path(file_path, expected_extensions=VALID_WORD_EXTS)
    bak_msg = ""
    if backup:
        bak = create_backup(path)
        if bak:
            bak_msg = f" (backup created: {bak.name})"

    doc = Document(str(path))
    if paragraph_idx < 0 or paragraph_idx >= len(doc.paragraphs):
        raise IndexError(
            f"Paragraph index {paragraph_idx} out of range. Document has {len(doc.paragraphs)} paragraphs."
        )

    p = doc.paragraphs[paragraph_idx]
    element = p._element
    element.getparent().remove(element)

    doc.save(str(path))
    return f"Deleted paragraph #{paragraph_idx} in {path}{bak_msg}"
