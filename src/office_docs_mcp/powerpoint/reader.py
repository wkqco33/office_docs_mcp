from __future__ import annotations

from typing import Any

from pptx import Presentation

from office_docs_mcp.common.file_utils import validate_file_path

VALID_PPT_EXTS = (".pptx",)


def ppt_get_outline(file_path: str) -> dict[str, Any]:
    """Get the outline of a PowerPoint presentation.

    Args:
        file_path: Path to the PowerPoint file (.pptx).

    Returns:
        Dictionary with total slide count and list of slides with their titles.
    """
    path = validate_file_path(file_path, expected_extensions=VALID_PPT_EXTS)
    prs = Presentation(str(path))

    slides_summary = []
    for idx, slide in enumerate(prs.slides):
        title_text = ""
        if slide.shapes.title and slide.shapes.title.has_text_frame:
            title_text = slide.shapes.title.text.strip()
        elif len(slide.shapes) > 0 and slide.shapes[0].has_text_frame:
            title_text = slide.shapes[0].text.strip()

        slides_summary.append(
            {
                "slide_index": idx,
                "title": title_text,
                "shapes_count": len(slide.shapes),
            }
        )

    return {
        "file_path": str(path),
        "total_slides": len(prs.slides),
        "slides": slides_summary,
    }


def ppt_read_slide(file_path: str, slide_idx: int) -> dict[str, Any]:
    """Read all text contents and tables from a specific slide.

    Args:
        file_path: Path to the PowerPoint file (.pptx).
        slide_idx: 0-based index of the slide.

    Returns:
        Dictionary containing slide details and shapes data.
    """
    path = validate_file_path(file_path, expected_extensions=VALID_PPT_EXTS)
    prs = Presentation(str(path))

    if slide_idx < 0 or slide_idx >= len(prs.slides):
        raise IndexError(
            f"Slide index {slide_idx} out of range. Presentation has {len(prs.slides)} slides."
        )

    slide = prs.slides[slide_idx]
    title_text = ""
    if slide.shapes.title and slide.shapes.title.has_text_frame:
        title_text = slide.shapes.title.text.strip()

    shapes_data = []
    for s_idx, shape in enumerate(slide.shapes):
        shape_info: dict[str, Any] = {
            "shape_index": s_idx,
            "name": shape.name,
            "has_text": shape.has_text_frame,
            "is_table": shape.has_table,
            "has_chart": shape.has_chart,
        }

        if shape.has_text_frame:
            shape_info["text"] = shape.text.strip()

        if shape.has_table:
            table = shape.table
            rows_data = []
            for row in table.rows:
                rows_data.append([cell.text.strip() for cell in row.cells])
            shape_info["table_data"] = rows_data

        if shape.has_chart:
            chart = shape.chart
            chart_title = ""
            if chart.has_title and chart.chart_title and chart.chart_title.has_text_frame:
                chart_title = chart.chart_title.text_frame.text.strip()
            shape_info["chart_title"] = chart_title
            shape_info["chart_type"] = str(chart.chart_type)

        shapes_data.append(shape_info)

    # Read notes if available
    notes_text = ""
    if slide.has_notes_slide and slide.notes_slide.notes_text_frame:
        notes_text = slide.notes_slide.notes_text_frame.text.strip()

    return {
        "slide_index": slide_idx,
        "title": title_text,
        "shapes": shapes_data,
        "notes": notes_text,
    }


def ppt_read_notes(file_path: str, slide_idx: int) -> str:
    """Read the speaker notes of a specific slide.

    Args:
        file_path: Path to the PowerPoint file (.pptx).
        slide_idx: 0-based index of the slide.

    Returns:
        Speaker notes text string, or empty string if no notes exist.
    """
    path = validate_file_path(file_path, expected_extensions=VALID_PPT_EXTS)
    prs = Presentation(str(path))

    if slide_idx < 0 or slide_idx >= len(prs.slides):
        raise IndexError(
            f"Slide index {slide_idx} out of range. Presentation has {len(prs.slides)} slides."
        )

    slide = prs.slides[slide_idx]
    if not slide.has_notes_slide or not slide.notes_slide.notes_text_frame:
        return ""

    return slide.notes_slide.notes_text_frame.text.strip()


def ppt_search(
    file_path: str,
    query: str,
    case_sensitive: bool = False,
    max_results: int = 50,
) -> list[dict[str, Any]]:
    """Search for a text query across PowerPoint slides, shapes, tables, and notes.

    Args:
        file_path: Path to the PowerPoint file (.pptx).
        query: Text to search for.
        case_sensitive: Whether search is case sensitive.
        max_results: Maximum results to return (default 50).

    Returns:
        List of match dictionaries detailing slide index, source, and text.
    """
    path = validate_file_path(file_path, expected_extensions=VALID_PPT_EXTS)
    prs = Presentation(str(path))
    results: list[dict[str, Any]] = []

    norm_query = query if case_sensitive else query.lower()

    for s_idx, slide in enumerate(prs.slides):
        slide_title = ""
        if slide.shapes.title and slide.shapes.title.has_text_frame:
            slide_title = slide.shapes.title.text.strip()

        # 1. Search shapes and tables
        for sh_idx, shape in enumerate(slide.shapes):
            if shape.has_text_frame:
                txt = shape.text
                cmp_txt = txt if case_sensitive else txt.lower()
                if norm_query in cmp_txt:
                    results.append(
                        {
                            "slide_idx": s_idx,
                            "slide_title": slide_title,
                            "source": "shape",
                            "shape_idx": sh_idx,
                            "shape_name": shape.name,
                            "text": txt.strip(),
                        }
                    )
                    if len(results) >= max_results:
                        return results

            if shape.has_table:
                for r_idx, row in enumerate(shape.table.rows):
                    for c_idx, cell in enumerate(row.cells):
                        cell_txt = cell.text
                        cmp_txt = cell_txt if case_sensitive else cell_txt.lower()
                        if norm_query in cmp_txt:
                            results.append(
                                {
                                    "slide_idx": s_idx,
                                    "slide_title": slide_title,
                                    "source": "table",
                                    "shape_idx": sh_idx,
                                    "row_idx": r_idx,
                                    "col_idx": c_idx,
                                    "text": cell_txt.strip(),
                                }
                            )
                            if len(results) >= max_results:
                                return results

        # 2. Search speaker notes
        if slide.has_notes_slide and slide.notes_slide.notes_text_frame:
            notes_txt = slide.notes_slide.notes_text_frame.text
            cmp_txt = notes_txt if case_sensitive else notes_txt.lower()
            if norm_query in cmp_txt:
                results.append(
                    {
                        "slide_idx": s_idx,
                        "slide_title": slide_title,
                        "source": "notes",
                        "text": notes_txt.strip(),
                    }
                )
                if len(results) >= max_results:
                    return results

    return results
