from __future__ import annotations

from typing import Any

from pptx import Presentation

from office_docs_mcp.common.file_utils import (
    create_backup,
    ensure_parent_dir,
    validate_file_path,
)

VALID_PPT_EXTS = (".pptx",)


def ppt_create_presentation(file_path: str, title: str | None = None) -> str:
    """Create a new PowerPoint presentation with an optional title slide.

    Args:
        file_path: Path where the presentation will be saved (.pptx).
        title: Optional title for the first title slide.

    Returns:
        Status confirmation message.
    """
    path = ensure_parent_dir(file_path)
    if path.suffix.lower() not in VALID_PPT_EXTS:
        path = path.with_suffix(".pptx")

    prs = Presentation()
    if title:
        # Layout 0 is typically the title slide layout
        title_slide_layout = prs.slide_layouts[0]
        slide = prs.slides.add_slide(title_slide_layout)
        if slide.shapes.title:
            slide.shapes.title.text = title

    prs.save(str(path))
    return f"Created PowerPoint presentation at {path}"


def ppt_add_slide(
    file_path: str,
    title: str,
    content: str | None = None,
    layout_idx: int = 1,
    backup: bool = False,
) -> str:
    """Add a new slide to a PowerPoint presentation.

    Args:
        file_path: Path to the PowerPoint file (.pptx).
        title: Title of the slide.
        content: Optional body text for the slide content placeholder.
        layout_idx: Slide layout index (0=Title, 1=Title and Content). Default 1.
        backup: If True, create a backup of the original file before modifying.

    Returns:
        Status confirmation message.
    """
    path = validate_file_path(file_path, expected_extensions=VALID_PPT_EXTS)
    bak_msg = ""
    if backup:
        bak = create_backup(path)
        if bak:
            bak_msg = f" (backup created: {bak.name})"

    prs = Presentation(str(path))

    if layout_idx < 0 or layout_idx >= len(prs.slide_layouts):
        layout_idx = 1

    slide_layout = prs.slide_layouts[layout_idx]
    slide = prs.slides.add_slide(slide_layout)

    if slide.shapes.title:
        slide.shapes.title.text = title

    if content:
        if len(slide.placeholders) > 1:
            body_shape = slide.placeholders[1]
            tf = body_shape.text_frame
            tf.text = content
        else:
            from pptx.util import Inches

            tb = slide.shapes.add_textbox(Inches(1), Inches(2), Inches(8), Inches(4))
            tb.text_frame.word_wrap = True
            tb.text_frame.text = content

    prs.save(str(path))
    new_idx = len(prs.slides) - 1
    return f"Added slide #{new_idx} with title '{title}' to {path}{bak_msg}"


def ppt_update_slide_text(
    file_path: str,
    slide_idx: int,
    shape_idx: int,
    text: str,
    backup: bool = False,
) -> str:
    """Update the text of a specific shape on a slide.

    Args:
        file_path: Path to the PowerPoint file (.pptx).
        slide_idx: 0-based slide index.
        shape_idx: 0-based shape index within the slide.
        text: New text content for the shape.
        backup: If True, create a backup of the original file before modifying.

    Returns:
        Status confirmation message.
    """
    path = validate_file_path(file_path, expected_extensions=VALID_PPT_EXTS)
    bak_msg = ""
    if backup:
        bak = create_backup(path)
        if bak:
            bak_msg = f" (backup created: {bak.name})"

    prs = Presentation(str(path))

    if slide_idx < 0 or slide_idx >= len(prs.slides):
        raise IndexError(
            f"Slide index {slide_idx} out of range. Presentation has {len(prs.slides)} slides."
        )

    slide = prs.slides[slide_idx]
    if shape_idx < 0 or shape_idx >= len(slide.shapes):
        raise IndexError(f"Shape index {shape_idx} out of range on slide {slide_idx}.")

    shape = slide.shapes[shape_idx]
    if not shape.has_text_frame:
        raise ValueError(f"Shape {shape_idx} does not support text editing.")

    shape.text_frame.text = text
    prs.save(str(path))
    return f"Updated text on slide {slide_idx}, shape {shape_idx} in {path}{bak_msg}"


def ppt_update_notes(
    file_path: str,
    slide_idx: int,
    notes_text: str,
    backup: bool = False,
) -> str:
    """Update the speaker notes of a specific slide.

    Args:
        file_path: Path to the PowerPoint file (.pptx).
        slide_idx: 0-based slide index.
        notes_text: Speaker notes text to set for the slide.
        backup: If True, create a backup of the original file before modifying.

    Returns:
        Status confirmation message.
    """
    path = validate_file_path(file_path, expected_extensions=VALID_PPT_EXTS)
    bak_msg = ""
    if backup:
        bak = create_backup(path)
        if bak:
            bak_msg = f" (backup created: {bak.name})"

    prs = Presentation(str(path))

    if slide_idx < 0 or slide_idx >= len(prs.slides):
        raise IndexError(
            f"Slide index {slide_idx} out of range. Presentation has {len(prs.slides)} slides."
        )

    slide = prs.slides[slide_idx]
    notes_slide = slide.notes_slide
    text_frame = notes_slide.notes_text_frame
    text_frame.text = notes_text

    prs.save(str(path))
    return f"Updated notes on slide {slide_idx} in {path}{bak_msg}"


def ppt_add_table(
    file_path: str,
    slide_idx: int,
    rows: int,
    cols: int,
    data: list[list[str]] | None = None,
    left: float = 1.0,
    top: float = 2.0,
    width: float = 8.0,
    height: float = 4.0,
    backup: bool = False,
) -> str:
    """Add a table to a specific slide in a PowerPoint presentation.

    Args:
        file_path: Path to the PowerPoint file (.pptx).
        slide_idx: 0-based index of the slide.
        rows: Number of table rows.
        cols: Number of table columns.
        data: Optional 2D list of cell text data.
        left: Left position in inches (default 1.0).
        top: Top position in inches (default 2.0).
        width: Table width in inches (default 8.0).
        height: Table height in inches (default 4.0).
        backup: If True, create a backup before modifying.

    Returns:
        Status confirmation message.
    """
    from pptx.util import Inches

    path = validate_file_path(file_path, expected_extensions=VALID_PPT_EXTS)
    bak_msg = ""
    if backup:
        bak = create_backup(path)
        if bak:
            bak_msg = f" (backup created: {bak.name})"

    prs = Presentation(str(path))
    if slide_idx < 0 or slide_idx >= len(prs.slides):
        raise IndexError(
            f"Slide index {slide_idx} out of range. Presentation has {len(prs.slides)} slides."
        )

    slide = prs.slides[slide_idx]
    shape = slide.shapes.add_table(
        rows, cols, Inches(left), Inches(top), Inches(width), Inches(height)
    )
    table = shape.table

    if data:
        for r_idx, row in enumerate(data):
            if r_idx < rows:
                for c_idx, val in enumerate(row):
                    if c_idx < cols:
                        table.cell(r_idx, c_idx).text = str(val)

    prs.save(str(path))
    return f"Added table ({rows}x{cols}) to slide {slide_idx} in {path}{bak_msg}"


def ppt_add_chart(
    file_path: str,
    slide_idx: int,
    chart_type: str,
    categories: list[str],
    series_data: list[dict[str, Any]],
    title: str | None = None,
    left: float = 1.0,
    top: float = 1.5,
    width: float = 8.0,
    height: float = 4.5,
    backup: bool = False,
) -> str:
    """Add a native Excel-backed chart to a slide in a PowerPoint presentation.

    Args:
        file_path: Path to the PowerPoint file (.pptx).
        slide_idx: 0-based index of the slide.
        chart_type: Type of chart ('column_clustered', 'bar_clustered', 'line', 'pie', 'doughnut', 'area', etc.).
        categories: List of category labels along the X-axis (e.g. ['Q1', 'Q2', 'Q3', 'Q4']).
        series_data: List of series dictionaries, e.g. [{'name': 'Sales', 'values': [100, 120, 140, 160]}].
        title: Optional chart title.
        left: Left offset position in inches (default: 1.0).
        top: Top offset position in inches (default: 1.5).
        width: Chart width in inches (default: 8.0).
        height: Chart height in inches (default: 4.5).
        backup: If True, create a backup before modifying.

    Returns:
        Status confirmation message.
    """
    from pptx.chart.data import CategoryChartData
    from pptx.enum.chart import XL_CHART_TYPE
    from pptx.util import Inches

    norm_chart_type = chart_type.lower().strip()
    chart_type_map = {
        "column_clustered": XL_CHART_TYPE.COLUMN_CLUSTERED,
        "column_stacked": XL_CHART_TYPE.COLUMN_STACKED,
        "column_stacked_100": XL_CHART_TYPE.COLUMN_STACKED_100,
        "bar_clustered": XL_CHART_TYPE.BAR_CLUSTERED,
        "bar_stacked": XL_CHART_TYPE.BAR_STACKED,
        "bar_stacked_100": XL_CHART_TYPE.BAR_STACKED_100,
        "line": XL_CHART_TYPE.LINE,
        "line_markers": XL_CHART_TYPE.LINE_MARKERS,
        "line_stacked": XL_CHART_TYPE.LINE_STACKED,
        "pie": XL_CHART_TYPE.PIE,
        "pie_exploded": XL_CHART_TYPE.PIE_EXPLODED,
        "doughnut": XL_CHART_TYPE.DOUGHNUT,
        "area": XL_CHART_TYPE.AREA,
        "area_stacked": XL_CHART_TYPE.AREA_STACKED,
    }
    if norm_chart_type not in chart_type_map:
        supported = ", ".join(chart_type_map.keys())
        raise ValueError(f"Unsupported chart type '{chart_type}'. Supported types: {supported}")

    path = validate_file_path(file_path, expected_extensions=VALID_PPT_EXTS)
    bak_msg = ""
    if backup:
        bak = create_backup(path)
        if bak:
            bak_msg = f" (backup created: {bak.name})"

    prs = Presentation(str(path))
    if slide_idx < 0 or slide_idx >= len(prs.slides):
        raise IndexError(
            f"Slide index {slide_idx} out of range. Presentation has {len(prs.slides)} slides."
        )

    slide = prs.slides[slide_idx]

    chart_data = CategoryChartData()
    chart_data.categories = categories
    for series in series_data:
        s_name = series.get("name", "")
        s_values = tuple(series.get("values", []))
        chart_data.add_series(s_name, s_values)

    xl_chart_type = chart_type_map[norm_chart_type]
    chart_shape = slide.shapes.add_chart(
        xl_chart_type,
        Inches(left),
        Inches(top),
        Inches(width),
        Inches(height),
        chart_data,
    )
    chart = chart_shape.chart

    if title:
        chart.has_title = True
        chart.chart_title.text_frame.text = title

    prs.save(str(path))
    return f"Added {norm_chart_type} chart to slide {slide_idx} in {path}{bak_msg}"


def ppt_add_flowchart(
    file_path: str,
    slide_idx: int,
    mermaid_code: str,
    title: str | None = None,
    direction: str = "auto",
    left: float = 0.8,
    top: float = 1.6,
    width: float = 8.4,
    height: float = 5.0,
    backup: bool = False,
) -> str:
    """Add an editable flowchart or architecture diagram to a slide using Mermaid syntax.

    Nodes are rendered as native PowerPoint shapes (Rounded Rectangles, Diamonds, etc.)
    connected by native connector arrows, making text, colors, and layout fully editable.

    Args:
        file_path: Path to the PowerPoint file (.pptx).
        slide_idx: 0-based index of the slide.
        mermaid_code: Mermaid syntax code (e.g. 'graph TD\n A[Client] --> B[Server]').
        title: Optional title to place at the top of the slide.
        direction: Layout direction: 'auto' (detect from code), 'TD' (Top-Down), or 'LR' (Left-Right).
        left: Left margin in inches (default: 0.8).
        top: Top margin in inches (default: 1.6).
        width: Total width allocated for the diagram in inches (default: 8.4).
        height: Total height allocated for the diagram in inches (default: 5.0).
        backup: If True, create a backup before modifying.

    Returns:
        Status confirmation message.
    """
    from office_docs_mcp.powerpoint.flowchart import parse_mermaid, render_flowchart

    path = validate_file_path(file_path, expected_extensions=VALID_PPT_EXTS)
    bak_msg = ""
    if backup:
        bak = create_backup(path)
        if bak:
            bak_msg = f" (backup created: {bak.name})"

    prs = Presentation(str(path))
    if len(prs.slides) == 0:
        # Automatically create a slide if presentation has no slides
        blank_layout = prs.slide_layouts[6] if len(prs.slide_layouts) > 6 else prs.slide_layouts[0]
        slide = prs.slides.add_slide(blank_layout)
        slide_idx = 0
    elif slide_idx < 0 or slide_idx >= len(prs.slides):
        raise IndexError(
            f"Slide index {slide_idx} out of range. Presentation has {len(prs.slides)} slides."
        )
    else:
        slide = prs.slides[slide_idx]

    nodes, edges, detected_dir = parse_mermaid(mermaid_code)
    actual_dir = detected_dir if direction == "auto" else direction.upper()

    render_flowchart(
        slide=slide,
        nodes=nodes,
        edges=edges,
        direction=actual_dir,
        left=left,
        top=top,
        width=width,
        height=height,
        title=title,
    )

    prs.save(str(path))
    return f"Added flowchart ({len(nodes)} nodes, {len(edges)} edges) to slide {slide_idx} in {path}{bak_msg}"
