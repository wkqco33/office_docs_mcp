from __future__ import annotations

from typing import Any

import openpyxl
from openpyxl.utils.cell import coordinate_to_tuple

from office_docs_mcp.common.file_utils import (
    create_backup,
    ensure_parent_dir,
    validate_file_path,
)

VALID_EXCEL_EXTS = (".xlsx", ".xlsm")


def excel_create_workbook(file_path: str, sheet_names: list[str] | None = None) -> str:
    """Create a new empty Excel workbook with specified sheets.

    Args:
        file_path: Path where the new Excel file will be saved.
        sheet_names: Optional list of sheet names to initialize. Defaults to ['Sheet'].

    Returns:
        Status confirmation message.
    """
    path = ensure_parent_dir(file_path)
    if path.suffix.lower() not in VALID_EXCEL_EXTS:
        path = path.with_suffix(".xlsx")

    wb = openpyxl.Workbook()
    sheets = sheet_names if sheet_names else ["Sheet"]

    # Setup sheets
    default_ws = wb.active
    default_ws.title = sheets[0]
    for s in sheets[1:]:
        wb.create_sheet(title=s)

    wb.save(path)
    wb.close()
    return f"Created Excel workbook at {path} with sheets: {sheets}"


def excel_write_cell(
    file_path: str,
    sheet_name: str,
    coordinate: str,
    value: Any,
    backup: bool = False,
) -> str:
    """Write a value to a specific cell (e.g. 'A1', 'C10').

    Args:
        file_path: Path to the Excel file.
        sheet_name: Target sheet name.
        coordinate: Cell coordinate (e.g. 'A1', 'B5').
        value: Value to write into the cell.
        backup: If True, create a backup before modifying.

    Returns:
        Status confirmation message.
    """
    path = validate_file_path(file_path, expected_extensions=VALID_EXCEL_EXTS)
    bak_msg = ""
    if backup:
        bak = create_backup(path)
        if bak:
            bak_msg = f" (backup created: {bak.name})"

    wb = openpyxl.load_workbook(path)

    try:
        if sheet_name not in wb.sheetnames:
            raise ValueError(f"Sheet '{sheet_name}' not found in workbook.")
        ws = wb[sheet_name]
        ws[coordinate] = value
        wb.save(path)
        return f"Written value to {sheet_name}!{coordinate} in {path}{bak_msg}"
    finally:
        wb.close()


def excel_write_range(
    file_path: str,
    sheet_name: str,
    start_cell: str,
    data: list[list[Any]],
    backup: bool = False,
) -> str:
    """Write a 2D matrix of data starting from a designated top-left cell.

    Args:
        file_path: Path to the Excel file.
        sheet_name: Target sheet name.
        start_cell: Top-left cell coordinate (e.g. 'A1').
        data: 2D array of values to write.
        backup: If True, create a backup before modifying.

    Returns:
        Status confirmation message.
    """
    path = validate_file_path(file_path, expected_extensions=VALID_EXCEL_EXTS)
    bak_msg = ""
    if backup:
        bak = create_backup(path)
        if bak:
            bak_msg = f" (backup created: {bak.name})"

    wb = openpyxl.load_workbook(path)

    try:
        if sheet_name not in wb.sheetnames:
            raise ValueError(f"Sheet '{sheet_name}' not found in workbook.")
        ws = wb[sheet_name]

        start_row, start_col = coordinate_to_tuple(start_cell)

        for r_idx, row in enumerate(data):
            for c_idx, val in enumerate(row):
                ws.cell(row=start_row + r_idx, column=start_col + c_idx, value=val)

        wb.save(path)
        rows_count = len(data)
        cols_count = max(len(r) for r in data) if data else 0
        return f"Written {rows_count} rows x {cols_count} cols starting at {sheet_name}!{start_cell}{bak_msg}"
    finally:
        wb.close()


def excel_append_rows(
    file_path: str,
    sheet_name: str,
    rows: list[list[Any]],
    backup: bool = False,
) -> str:
    """Append multiple rows to the end of a sheet.

    Args:
        file_path: Path to the Excel file.
        sheet_name: Target sheet name.
        rows: List of row values to append.
        backup: If True, create a backup before modifying.

    Returns:
        Status confirmation message.
    """
    path = validate_file_path(file_path, expected_extensions=VALID_EXCEL_EXTS)
    bak_msg = ""
    if backup:
        bak = create_backup(path)
        if bak:
            bak_msg = f" (backup created: {bak.name})"

    wb = openpyxl.load_workbook(path)

    try:
        if sheet_name not in wb.sheetnames:
            raise ValueError(f"Sheet '{sheet_name}' not found in workbook.")
        ws = wb[sheet_name]

        for row in rows:
            ws.append(row)

        wb.save(path)
        return f"Appended {len(rows)} rows to {sheet_name} in {path}{bak_msg}"
    finally:
        wb.close()


def excel_manage_sheets(
    file_path: str,
    action: str,
    sheet_name: str,
    new_name: str | None = None,
    backup: bool = False,
) -> str:
    """Manage sheets in an Excel workbook (add, rename, delete, copy).

    Args:
        file_path: Path to the Excel file.
        action: Action to perform: 'add', 'rename', 'delete', or 'copy'.
        sheet_name: Target sheet name.
        new_name: Required for 'rename' and 'copy' actions.
        backup: If True, create a backup before modifying.

    Returns:
        Status confirmation message.
    """
    action_norm = action.lower().strip()
    if action_norm not in ("add", "rename", "delete", "copy"):
        raise ValueError(
            f"Unsupported action '{action}'. Supported actions: add, rename, delete, copy"
        )

    path = validate_file_path(file_path, expected_extensions=VALID_EXCEL_EXTS)
    bak_msg = ""
    if backup:
        bak = create_backup(path)
        if bak:
            bak_msg = f" (backup created: {bak.name})"

    wb = openpyxl.load_workbook(path)

    try:
        if action_norm == "add":
            if sheet_name in wb.sheetnames:
                raise ValueError(f"Sheet '{sheet_name}' already exists.")
            wb.create_sheet(title=sheet_name)
            wb.save(path)
            return f"Added sheet '{sheet_name}' to {path}{bak_msg}"

        if sheet_name not in wb.sheetnames:
            raise ValueError(f"Sheet '{sheet_name}' not found in workbook.")

        if action_norm == "rename":
            if not new_name:
                raise ValueError("new_name is required for rename action.")
            if new_name in wb.sheetnames:
                raise ValueError(f"Sheet '{new_name}' already exists.")
            wb[sheet_name].title = new_name
            wb.save(path)
            return f"Renamed sheet '{sheet_name}' to '{new_name}' in {path}{bak_msg}"

        if action_norm == "copy":
            if not new_name:
                raise ValueError("new_name is required for copy action.")
            source = wb[sheet_name]
            target = wb.copy_worksheet(source)
            target.title = new_name
            wb.save(path)
            return f"Copied sheet '{sheet_name}' to '{new_name}' in {path}{bak_msg}"

        if action_norm == "delete":
            if len(wb.sheetnames) <= 1:
                raise ValueError("Cannot delete the only sheet in a workbook.")
            del wb[sheet_name]
            wb.save(path)
            return f"Deleted sheet '{sheet_name}' from {path}{bak_msg}"

        raise ValueError(f"Unsupported action '{action}'.")
    finally:
        wb.close()
