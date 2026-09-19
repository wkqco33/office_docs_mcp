from __future__ import annotations

from typing import Any

import openpyxl
from openpyxl.utils import get_column_letter

from office_docs_mcp.common.file_utils import format_as_markdown_table, validate_file_path

VALID_EXCEL_EXTS = (".xlsx", ".xlsm", ".xltx", ".xltm")


def excel_get_metadata(file_path: str) -> dict[str, Any]:
    """Get metadata of an Excel workbook including sheet names and dimensions.

    Args:
        file_path: Path to the Excel file (.xlsx, .xlsm).

    Returns:
        Dictionary containing sheet names, active sheet, and sheet summary.
    """
    path = validate_file_path(file_path, expected_extensions=VALID_EXCEL_EXTS)
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        sheet_summaries = []
        for name in wb.sheetnames:
            ws = wb[name]
            # Fetch first row if exists for column summary
            first_row = []
            if ws.max_row and ws.max_row > 0:
                for row in ws.iter_rows(min_row=1, max_row=1, values_only=True):
                    first_row = [str(c) if c is not None else "" for c in row]
                    break

            sheet_summaries.append(
                {
                    "name": name,
                    "max_row": ws.max_row,
                    "max_column": ws.max_column,
                    "headers_preview": first_row[:20] if first_row else [],
                }
            )

        active_title = (
            wb.active.title if wb.active else (wb.sheetnames[0] if wb.sheetnames else None)
        )
        return {
            "file_path": str(path),
            "sheet_names": wb.sheetnames,
            "active_sheet": active_title,
            "sheets": sheet_summaries,
        }
    finally:
        wb.close()


def excel_read_sheet(
    file_path: str,
    sheet_name: str | None = None,
    start_row: int = 1,
    end_row: int = 50,
    start_col: int = 1,
    end_col: int = 20,
    evaluate_formulas: bool = True,
    format: str = "markdown",
) -> str | list[list[Any]]:
    """Read a rectangular cell range from an Excel sheet.

    Indices are 1-based: row 1 is the first row, col 1 is column 'A'.

    Args:
        file_path: Path to the Excel file.
        sheet_name: Target sheet name (defaults to active sheet).
        start_row: First row to read (1-based, inclusive). Default 1.
        end_row: Last row to read (1-based, inclusive). Default 50.
        start_col: First column to read (1-based, inclusive). Default 1.
        end_col: Last column to read (1-based, inclusive). Default 20.
        evaluate_formulas: If True, reads calculated values. If False, reads raw formulas.
        format: Output format, either 'markdown' (default) or 'raw' (list of lists).

    Returns:
        Markdown table string or 2D list of values.
    """
    path = validate_file_path(file_path, expected_extensions=VALID_EXCEL_EXTS)
    wb = openpyxl.load_workbook(path, data_only=evaluate_formulas)

    try:
        if sheet_name:
            if sheet_name not in wb.sheetnames:
                raise ValueError(
                    f"Sheet '{sheet_name}' not found. Available sheets: {wb.sheetnames}"
                )
            ws = wb[sheet_name]
        else:
            ws = wb.active

        # Clamp row and col boundaries
        actual_start_row = max(1, start_row)
        actual_end_row = max(actual_start_row, end_row)
        actual_start_col = max(1, start_col)
        actual_end_col = max(actual_start_col, end_col)

        rows_data = [
            list(row)
            for row in ws.iter_rows(
                min_row=actual_start_row,
                max_row=actual_end_row,
                min_col=actual_start_col,
                max_col=actual_end_col,
                values_only=True,
            )
        ]

        if format == "raw":
            return rows_data

        # Default markdown format
        headers = [get_column_letter(c) for c in range(actual_start_col, actual_end_col + 1)]
        return format_as_markdown_table(headers, rows_data)

    finally:
        wb.close()


def excel_search(
    file_path: str,
    query: str,
    sheet_name: str | None = None,
    case_sensitive: bool = False,
    max_results: int = 50,
) -> list[dict[str, Any]]:
    """Search for a text query across Excel worksheet cells.

    Args:
        file_path: Path to the Excel file.
        query: Text to search for.
        sheet_name: Specific sheet to search in. If None, searches all sheets.
        case_sensitive: Whether the search should be case sensitive.
        max_results: Maximum number of search results to return (default 50).

    Returns:
        List of dictionaries with sheet, coordinate, row, col, value, and row_preview.
    """
    path = validate_file_path(file_path, expected_extensions=VALID_EXCEL_EXTS)
    wb = openpyxl.load_workbook(path, data_only=True)
    results: list[dict[str, Any]] = []

    target_sheets = [sheet_name] if sheet_name else wb.sheetnames
    norm_query = query if case_sensitive else query.lower()

    try:
        for s_name in target_sheets:
            if s_name not in wb.sheetnames:
                continue
            ws = wb[s_name]

            for row in ws.iter_rows(values_only=False):
                row_vals = [cell.value for cell in row]
                for cell in row:
                    val_str = str(cell.value) if cell.value is not None else ""
                    cmp_str = val_str if case_sensitive else val_str.lower()
                    if norm_query in cmp_str:
                        results.append(
                            {
                                "sheet": s_name,
                                "coordinate": cell.coordinate,
                                "row": cell.row,
                                "col": cell.column,
                                "value": cell.value,
                                "row_preview": [v for v in row_vals if v is not None],
                            }
                        )
                        if len(results) >= max_results:
                            return results
        return results
    finally:
        wb.close()
