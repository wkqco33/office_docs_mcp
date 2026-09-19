from __future__ import annotations

import pytest

from office_docs_mcp.word.reader import (
    word_get_outline,
    word_read_paragraphs,
    word_read_table,
)
from office_docs_mcp.word.writer import (
    word_append_paragraph,
    word_append_table_row,
    word_create_document,
    word_write_table_cell,
)


def test_word_create_and_outline(tmp_path):
    f = tmp_path / "sample.docx"
    res = word_create_document(str(f), title="Project Overview")
    assert "Created" in res
    assert f.exists()

    word_append_paragraph(str(f), "Section 1: Background", style="Heading 1")
    word_append_paragraph(str(f), "This is the first paragraph describing the project.")

    outline = word_get_outline(str(f))
    assert outline["total_paragraphs"] >= 3
    assert any(h["text"] == "Project Overview" for h in outline["headings"])
    assert any(h["text"] == "Section 1: Background" for h in outline["headings"])


def test_word_read_paragraphs_slicing(tmp_path):
    f = tmp_path / "paragraphs.docx"
    word_create_document(str(f))
    for i in range(10):
        word_append_paragraph(str(f), f"Paragraph number {i}")

    # Read paragraphs 2 to 5 (count=4)
    paras = word_read_paragraphs(str(f), start_idx=2, count=4)
    assert len(paras) == 4
    assert paras[0]["text"] == "Paragraph number 2"
    assert paras[3]["text"] == "Paragraph number 5"


def test_word_table_operations(tmp_path):
    f = tmp_path / "table_doc.docx"
    word_create_document(str(f))

    # Add paragraph
    word_append_paragraph(str(f), "Table Section", style="Heading 1")

    # Add a table with initial row
    from docx import Document

    doc = Document(str(f))
    tbl = doc.add_table(rows=1, cols=2)
    tbl.rows[0].cells[0].text = "Item"
    tbl.rows[0].cells[1].text = "Price"
    doc.save(str(f))

    # Append row
    word_append_table_row(str(f), table_idx=0, row_data=["Apple", "$1.00"])

    # Check outline table count
    outline = word_get_outline(str(f))
    assert outline["total_tables"] == 1

    # Read table as markdown
    tbl_md = word_read_table(str(f), table_idx=0, format="markdown")
    assert "Item" in tbl_md
    assert "Apple" in tbl_md
    assert "$1.00" in tbl_md

    # Update cell
    word_write_table_cell(str(f), table_idx=0, row_idx=1, col_idx=1, text="$1.50")
    tbl_raw = word_read_table(str(f), table_idx=0, format="raw")
    assert tbl_raw[1][1] == "$1.50"


def test_word_invalid_table_index(tmp_path):
    f = tmp_path / "empty.docx"
    word_create_document(str(f))
    with pytest.raises(IndexError, match="Table index 0 out of range"):
        word_read_table(str(f), table_idx=0)


def test_word_invalid_style(tmp_path):
    f = tmp_path / "style_doc.docx"
    word_create_document(str(f))
    with pytest.raises(ValueError, match="Style 'NonExistentStyle' not found"):
        word_append_paragraph(str(f), "Text", style="NonExistentStyle")


def test_word_append_backup(tmp_path):
    f = tmp_path / "doc_bak.docx"
    word_create_document(str(f), title="Original")
    res = word_append_paragraph(str(f), "New text", backup=True)
    assert "backup created" in res
    backups = list(tmp_path.glob("doc_bak.docx.*.bak"))
    assert len(backups) == 1


def test_word_search(tmp_path):
    from docx import Document

    from office_docs_mcp.word.reader import word_search

    f = tmp_path / "search_doc.docx"
    word_create_document(str(f), title="Annual AI Strategy")
    word_append_paragraph(str(f), "The mission of our AI unit is innovation.", style="Heading 1")
    word_append_paragraph(str(f), "We develop agentic workflows for enterprise customers.")

    # Add a table with AI mentioned
    doc = Document(str(f))
    tbl = doc.add_table(rows=2, cols=2)
    tbl.rows[0].cells[0].text = "Project"
    tbl.rows[0].cells[1].text = "Scope"
    tbl.rows[1].cells[0].text = "Antigravity AI"
    tbl.rows[1].cells[1].text = "Coding Assistance"
    doc.save(str(f))

    # Case-insensitive search
    results = word_search(str(f), query="ai")
    assert len(results) >= 3  # In title, heading, and table cell
    types = {r["location_type"] for r in results}
    assert "paragraph" in types
    assert "table" in types

    # Case-sensitive search
    exact_results = word_search(str(f), query="Antigravity AI", case_sensitive=True)
    assert len(exact_results) == 1
    assert exact_results[0]["location_type"] == "table"
    assert exact_results[0]["table_idx"] == 0


def test_word_replace_text(tmp_path):
    from office_docs_mcp.word.reader import word_read_paragraphs
    from office_docs_mcp.word.writer import word_replace_text

    f = tmp_path / "template.docx"
    word_create_document(str(f), title="Contract for {{CLIENT_NAME}}")
    word_append_paragraph(str(f), "Agreement entered on {{DATE}} with {{CLIENT_NAME}}.")

    res = word_replace_text(
        str(f), find_text="{{CLIENT_NAME}}", replace_text="Acme Corp", backup=True
    )
    assert "Replaced 2 occurrence(s)" in res

    paras = word_read_paragraphs(str(f), start_idx=0, count=5)
    assert "Acme Corp" in paras[0]["text"]
    assert "Acme Corp" in paras[1]["text"]
    assert "{{CLIENT_NAME}}" not in paras[0]["text"]


def test_word_delete_paragraph(tmp_path):
    from office_docs_mcp.word.reader import word_get_outline, word_read_paragraphs
    from office_docs_mcp.word.writer import word_delete_paragraph

    f = tmp_path / "delete_doc.docx"
    word_create_document(str(f))
    word_append_paragraph(str(f), "Keep 1")
    word_append_paragraph(str(f), "Delete Me")
    word_append_paragraph(str(f), "Keep 2")

    initial_count = word_get_outline(str(f))["total_paragraphs"]

    res = word_delete_paragraph(str(f), paragraph_idx=1)
    assert "Deleted paragraph #1" in res

    new_count = word_get_outline(str(f))["total_paragraphs"]
    assert new_count == initial_count - 1

    paras = word_read_paragraphs(str(f), start_idx=0, count=5)
    texts = [p["text"] for p in paras]
    assert "Delete Me" not in texts
    assert "Keep 1" in texts
    assert "Keep 2" in texts
